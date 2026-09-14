import streamlit as st
from docx import Document
import re
import random
import hashlib

# ============================================================
# MCQ Master - Word Document Quiz App
# ============================================================

st.set_page_config(
    page_title="MCQ Master | Online Quiz",
    page_icon="📝",
    layout="wide",
    initial_sidebar_state="expanded",
)

# -------------------- Styling --------------------
st.markdown(
    """
    <style>
        .stApp {
            background: linear-gradient(135deg, #f5f7ff 0%, #eef2ff 100%);
        }

        .hero {
            padding: 30px;
            border-radius: 22px;
            background: linear-gradient(135deg, #4f46e5, #7c3aed);
            color: white;
            margin-bottom: 24px;
            box-shadow: 0 12px 30px rgba(79,70,229,.20);
        }

        .hero h1 {
            margin: 0;
            font-size: 2.5rem;
            font-weight: 800;
        }

        .hero p {
            margin: 8px 0 0;
            font-size: 1.05rem;
            opacity: .94;
        }

        .stat {
            padding: 18px 12px;
            border-radius: 16px;
            background: white;
            border: 1px solid #e5e7eb;
            text-align: center;
            box-shadow: 0 4px 15px rgba(0,0,0,.05);
        }

        .stat-number {
            font-size: 2rem;
            font-weight: 800;
            color: #4f46e5;
        }

        .stat-label {
            color: #6b7280;
            font-size: .9rem;
        }

        .question-card {
            padding: 20px;
            margin: 14px 0 8px;
            border-radius: 16px;
            background: white;
            border-left: 5px solid #6366f1;
            box-shadow: 0 4px 15px rgba(0,0,0,.05);
        }

        .result-card {
            padding: 28px;
            border-radius: 20px;
            background: white;
            text-align: center;
            box-shadow: 0 8px 25px rgba(0,0,0,.08);
        }

        div.stButton > button,
        div[data-testid="stFormSubmitButton"] button {
            border-radius: 10px;
            font-weight: 700;
        }

        .footer {
            text-align: center;
            color: #6b7280;
            padding: 15px;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

# -------------------- Session State --------------------
DEFAULTS = {
    "questions": [],
    "submitted": False,
    "answers": {},
    "file_id": None,
    "quiz_order": [],
}

for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = value.copy() if isinstance(value, (dict, list)) else value


# -------------------- Text Utilities --------------------
def clean_text(text: str) -> str:
    """Normalize whitespace while preserving the actual question text."""
    return re.sub(r"\s+", " ", str(text)).strip()


def strip_duplicate_question_number(text: str) -> str:
    """
    Handles documents where the question number is accidentally repeated,
    e.g. '1. 1. What is MATLAB?' -> 'What is MATLAB?'
    """
    text = clean_text(text)
    pattern = r"^(\d+)\s*[\.\):\-]\s*(?:\1)\s*[\.\):\-]\s*(.+)$"
    match = re.match(pattern, text)
    return match.group(2).strip() if match else text


def normalize_answer(answer):
    """
    Convert common answer formats to A/B/C/D.
    Examples:
      B
      B. Matrix Laboratory
      (B) Matrix Laboratory
      Option B
      Choice B
    """
    if not answer:
        return None

    answer = clean_text(answer)

    match = re.match(r"^[\(\[]?([A-Da-d])[\)\].:\-]?(?:\s|$)", answer)
    if match:
        return match.group(1).upper()

    match = re.search(r"\b(?:option|choice)\s*([A-Da-d])\b", answer, re.I)
    if match:
        return match.group(1).upper()

    return None


# -------------------- DOCX Parser --------------------
def extract_docx_lines(uploaded_file):
    """Read paragraphs and table cells from a DOCX file."""
    doc = Document(uploaded_file)
    lines = []

    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if text:
            lines.append(text)

    # Support MCQs stored inside Word tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = clean_text(cell.text)
                if text:
                    lines.append(text)

    return lines


def parse_docx(uploaded_file):
    """
    Parse MCQs in this format:

    1. Question text
    A. Option A
    B. Option B
    C. Option C
    D. Option D
    Answer: C
    """
    lines = extract_docx_lines(uploaded_file)

    questions = []
    current = None
    last_option = None

    def save_current():
        nonlocal current, last_option

        if not current:
            return

        # Only accept MCQs with A-D options.
        if set(current["options"].keys()) >= {"A", "B", "C", "D"}:
            answer = normalize_answer(current.get("answer"))

            if answer in current["options"]:
                questions.append(
                    {
                        "question": clean_text(current["question"]),
                        "options": {
                            letter: clean_text(current["options"][letter])
                            for letter in ["A", "B", "C", "D"]
                        },
                        "answer": answer,
                    }
                )

        current = None
        last_option = None

    for raw_line in lines:
        line = clean_text(raw_line)

        # Question formats:
        # 1. Question
        # 1) Question
        # Q1. Question
        # Q. Question
        # Question 1: Question
        q_match = re.match(
            r"^(?:Q(?:uestion)?\s*)?(\d+)\s*[\.\):\-]\s*(.+)$",
            line,
            re.IGNORECASE,
        )

        question_word_match = re.match(
            r"^Question\s+(\d+)\s*[\.\):\-]\s*(.+)$",
            line,
            re.IGNORECASE,
        )

        q_short = re.match(
            r"^Q\s*[\.\:\-]\s*(.+)$",
            line,
            re.IGNORECASE,
        )

        if q_match or question_word_match or q_short:
            save_current()

            if q_match:
                question_text = q_match.group(2)
            elif question_word_match:
                question_text = question_word_match.group(2)
            else:
                question_text = q_short.group(1)

            current = {
                "question": strip_duplicate_question_number(question_text),
                "options": {},
                "answer": None,
            }
            last_option = None
            continue

        # Option formats: A. / A) / (A) / A:
        option_match = re.match(
            r"^\(?([A-Da-d])\)?\s*[\.\:\)\-]\s*(.+)$",
            line,
        )

        if option_match and current:
            letter = option_match.group(1).upper()
            current["options"][letter] = option_match.group(2).strip()
            last_option = letter
            continue

        # Answer formats:
        # Answer: C
        # Correct Answer: C. ...
        # Ans: C
        answer_match = re.match(
            r"^(?:Correct\s+Answer|Correct\s+Option|Answer|Ans|Key)"
            r"\s*[\:\-\=]\s*(.+)$",
            line,
            re.IGNORECASE,
        )

        if answer_match and current:
            current["answer"] = answer_match.group(1).strip()
            continue

        # Continuation line.
        if current:
            if last_option and last_option in current["options"]:
                current["options"][last_option] += " " + line
            else:
                current["question"] += " " + line

    save_current()

    # Remove duplicate questions.
    unique_questions = []
    seen = set()

    for question in questions:
        key = re.sub(r"[^a-z0-9]+", " ", question["question"].lower()).strip()
        if key and key not in seen:
            seen.add(key)
            unique_questions.append(question)

    return unique_questions


# -------------------- Quiz Functions --------------------
def reset_quiz():
    st.session_state.submitted = False
    st.session_state.answers = {}

    # Clear radio widget values from the previous quiz.
    for key in list(st.session_state.keys()):
        if key.startswith("answer_"):
            del st.session_state[key]


def calculate_result(questions):
    correct = 0
    wrong = 0
    unanswered = 0

    for index, question in enumerate(questions):
        selected = st.session_state.answers.get(index)

        if selected is None:
            unanswered += 1
        elif selected == question["answer"]:
            correct += 1
        else:
            wrong += 1

    total = len(questions)
    percentage = (correct / total * 100) if total else 0

    return correct, wrong, unanswered, percentage


# -------------------- Header --------------------
st.markdown(
    """
    <div class="hero">
        <h1>📝 MCQ Master</h1>
        <p>Upload a Word document, attempt the MCQs, submit your answers, and instantly view your result.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

# -------------------- Sidebar --------------------
with st.sidebar:
    st.header("⚙️ Quiz Settings")

    uploaded_file = st.file_uploader(
        "Upload MCQ Word File",
        type=["docx"],
        help="Upload a .docx file containing questions, A-D options and correct answers.",
    )

    shuffle_questions = st.checkbox(
        "🔀 Shuffle questions",
        value=False,
    )

    show_result_details = st.checkbox(
        "📊 Show answer review",
        value=True,
    )

    st.markdown("---")
    st.subheader("📄 Supported Format")
    st.code(
        "1. What is MATLAB?\n"
        "A. Mathematical Laboratory\n"
        "B. Matrix Laboratory\n"
        "C. Machine Laboratory\n"
        "D. Matrix Language\n"
        "Answer: B",
        language="text",
    )

    st.caption("The app supports MCQs stored in Word paragraphs or tables.")


# -------------------- File Processing --------------------
if uploaded_file is not None:
    file_bytes = uploaded_file.getvalue()
    file_id = hashlib.md5(file_bytes).hexdigest()

    # Re-parse only when the uploaded document changes.
    if st.session_state.file_id != file_id:
        try:
            parsed_questions = parse_docx(uploaded_file)

            if shuffle_questions:
                random.shuffle(parsed_questions)

            st.session_state.questions = parsed_questions
            st.session_state.file_id = file_id
            st.session_state.quiz_order = list(range(len(parsed_questions)))
            reset_quiz()

        except Exception as exc:
            st.error(f"Unable to read the Word document: {exc}")
            st.stop()

    questions = st.session_state.questions

    if not questions:
        st.error(
            "No valid MCQs were detected. Please make sure each question has "
            "four options (A-D) and an answer line such as 'Answer: B'."
        )
        st.stop()

    # Re-order questions when the checkbox changes.
    if shuffle_questions and not st.session_state.get("shuffle_applied", False):
        random.shuffle(questions)
        st.session_state.questions = questions
        st.session_state.shuffle_applied = True
        reset_quiz()
        st.rerun()

    if not shuffle_questions:
        st.session_state.shuffle_applied = False

    total = len(questions)

    # -------------------- Statistics --------------------
    answered = len(st.session_state.answers)
    remaining = max(total - answered, 0)

    c1, c2, c3 = st.columns(3)

    with c1:
        st.markdown(
            f"""
            <div class="stat">
                <div class="stat-number">{total}</div>
                <div class="stat-label">Total Questions</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with c2:
        st.markdown(
            f"""
            <div class="stat">
                <div class="stat-number">{answered}</div>
                <div class="stat-label">Answered</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with c3:
        st.markdown(
            f"""
            <div class="stat">
                <div class="stat-number">{remaining}</div>
                <div class="stat-label">Remaining</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown("### 🎯 Student Quiz")

    # -------------------- Quiz Form --------------------
    with st.form("mcq_form", clear_on_submit=False):
        current_answers = {}

        for index, question in enumerate(questions):
            st.markdown(
                f"""
                <div class="question-card">
                    <strong>Question {index + 1}</strong>
                    <br><br>
                    {question["question"]}
                </div>
                """,
                unsafe_allow_html=True,
            )

            options = [
                f"{letter}. {question['options'][letter]}"
                for letter in ["A", "B", "C", "D"]
            ]

            selected = st.radio(
                "Select your answer:",
                options=options,
                index=None,
                key=f"answer_{index}",
                label_visibility="collapsed",
            )

            if selected:
                current_answers[index] = selected[0]

        submitted = st.form_submit_button(
            "🚀 Submit Quiz",
            use_container_width=True,
            type="primary",
        )

    # Save all selections only after the form is submitted.
    if submitted:
        st.session_state.answers = current_answers
        st.session_state.submitted = True

    # -------------------- Result --------------------
    if st.session_state.submitted:
        correct, wrong, unanswered, percentage = calculate_result(questions)

        st.markdown("---")
        st.markdown("## 🏆 Quiz Result")

        r1, r2, r3, r4 = st.columns(4)

        with r1:
            st.markdown(
                f'<div class="stat"><div class="stat-number">{correct}</div>'
                '<div class="stat-label">✅ Correct</div></div>',
                unsafe_allow_html=True,
            )

        with r2:
            st.markdown(
                f'<div class="stat"><div class="stat-number">{wrong}</div>'
                '<div class="stat-label">❌ Wrong</div></div>',
                unsafe_allow_html=True,
            )

        with r3:
            st.markdown(
                f'<div class="stat"><div class="stat-number">{unanswered}</div>'
                '<div class="stat-label">⏭️ Unanswered</div></div>',
                unsafe_allow_html=True,
            )

        with r4:
            st.markdown(
                f'<div class="stat"><div class="stat-number">{percentage:.1f}%</div>'
                '<div class="stat-label">📈 Score</div></div>',
                unsafe_allow_html=True,
            )

        st.progress(min(percentage / 100, 1.0))

        if percentage >= 80:
            st.success("🌟 Excellent performance! Keep it up.")
        elif percentage >= 60:
            st.info("👍 Good performance. A little more practice will help.")
        elif percentage >= 40:
            st.warning("📚 Keep practicing and review the incorrect answers.")
        else:
            st.error("💪 Keep practicing and review the incorrect answers.")

        if show_result_details:
            st.markdown("### 🔍 Answer Review")

            for index, question in enumerate(questions):
                selected = st.session_state.answers.get(index)
                correct_answer = question["answer"]

                if selected == correct_answer:
                    st.success(
                        f"**Q{index + 1}. Correct** — "
                        f"Your answer: {selected}. "
                        f"{question['options'][selected]}"
                    )
                elif selected:
                    st.error(
                        f"**Q{index + 1}. Wrong** — "
                        f"Your answer: {selected}. {question['options'][selected]}  \n"
                        f"**Correct:** {correct_answer}. "
                        f"{question['options'][correct_answer]}"
                    )
                else:
                    st.warning(
                        f"**Q{index + 1}. Not answered** — "
                        f"Correct answer: {correct_answer}. "
                        f"{question['options'][correct_answer]}"
                    )

        if st.button("🔄 Retake Quiz", use_container_width=True):
            reset_quiz()
            st.rerun()

else:
    st.markdown(
        """
        <div class="result-card">
            <h2>👋 Welcome to MCQ Master</h2>
            <p>Upload your Word document from the sidebar to start the quiz.</p>
            <p><b>Supported:</b> .docx files containing MCQs</p>
            <p><b>Workflow:</b> Upload → Attempt → Submit → View Score</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.markdown(
    '<div class="footer">MCQ Master • Built with Streamlit • Word-to-Quiz Learning Tool</div>',
    unsafe_allow_html=True,
)
